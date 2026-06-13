from __future__ import annotations

import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "vectra_x_outputs"
TABLES = DATA / "tables"
OUT = ROOT / "outputs" / "submission"
ASSETS = OUT / "assets"
DOCX_PATH = OUT / "VECTRA_X_Technical_Report.docx"
PDF_PATH = OUT / "VECTRA_X_Technical_Report.pdf"

NAVY = RGBColor(8, 31, 48)
TEAL = RGBColor(28, 151, 137)
BLUE = RGBColor(37, 93, 125)
GRAY = RGBColor(89, 103, 113)
LIGHT = "EAF1F4"
PALE_TEAL = "E7F5F2"
PALE_RED = "FBECEB"
WHITE = RGBColor(255, 255, 255)


def load(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES / name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=11, bold=False, color=None, italic=False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_paragraph(doc, text: str, *, bold_prefix: str | None = None, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_font(rest, italic=italic)
    else:
        run = paragraph.add_run(text)
        set_font(run, italic=italic)
    return paragraph


def add_bullet(doc, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.194)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_font(
        run,
        size={1: 16, 2: 13, 3: 12}[level],
        bold=True,
        color=BLUE if level < 3 else NAVY,
    )
    return paragraph


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_font(run, size=9, bold=True, color=NAVY)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(str(value))
            set_font(run, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, title: str, text: str, fill=PALE_TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_cell_shading(table.cell(0, 0), fill)
    paragraph = table.cell(0, 0).paragraphs[0]
    title_run = paragraph.add_run(f"{title}\n")
    set_font(title_run, size=10.5, bold=True, color=NAVY)
    text_run = paragraph.add_run(text)
    set_font(text_run, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def make_figures(model_comparison: pd.DataFrame, per_label: pd.DataFrame) -> tuple[Path, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    model_path = ASSETS / "model_comparison.png"
    label_path = ASSETS / "per_label_recall.png"

    fig, ax = plt.subplots(figsize=(8.2, 4.5))
    view = model_comparison.sort_values("macro_pr_auc")
    ax.barh(view["model"], view["macro_pr_auc"], color="#1C9789")
    ax.set_xlabel("Training-only OOF macro PR-AUC")
    ax.set_title("Model ranking under label imbalance")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(model_path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    colors = ["#1C9789" if value >= 0.5 else "#C6534C" for value in per_label["recall"]]
    ax.bar(per_label["label"], per_label["recall"], color=colors)
    ax.set_ylim(0, 1)
    ax.set_ylabel("Held-out recall")
    ax.set_title("Rare-label false-negative risk remains the key limitation")
    ax.spines[["top", "right"]].set_visible(False)
    for index, row in per_label.reset_index(drop=True).iterrows():
        ax.text(index, row["recall"] + 0.025, f"n={int(row['support_pos'])}", ha="center", fontsize=8)
    fig.tight_layout()
    fig.savefig(label_path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return model_path, label_path


def build_pdf(
    manifest,
    models,
    heldout,
    per_label,
    tracks,
    confidence,
    calibration,
    conformal,
    transfer,
    resources,
    thresholds,
    model_figure,
    label_figure,
) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "VXBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10,
        leading=13.2, spaceAfter=8, textColor=colors.HexColor("#182A35"),
    )
    h1 = ParagraphStyle(
        "VXH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16,
        leading=19, spaceBefore=13, spaceAfter=7, textColor=colors.HexColor("#255D7D"),
    )
    h2 = ParagraphStyle(
        "VXH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12,
        leading=15, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#08304A"),
    )
    cover_title = ParagraphStyle(
        "CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=30,
        leading=34, alignment=TA_CENTER, textColor=colors.HexColor("#081F30"), spaceAfter=10,
    )
    cover_subtitle = ParagraphStyle(
        "CoverSubtitle", parent=body, fontName="Helvetica", fontSize=15, leading=19,
        alignment=TA_CENTER, textColor=colors.HexColor("#255D7D"), spaceAfter=18,
    )
    caption = ParagraphStyle(
        "Caption", parent=body, fontName="Helvetica-Oblique", fontSize=8.5,
        leading=10, alignment=TA_CENTER, textColor=colors.HexColor("#596771"), spaceAfter=8,
    )
    bullet = ParagraphStyle(
        "Bullet", parent=body, leftIndent=18, firstLineIndent=-9, bulletIndent=6, spaceAfter=4,
    )

    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#596771"))
        canvas.drawString(0.8 * inch, 10.55 * inch, "VECTRA-X | FIT Competition 2026, Track IV")
        canvas.drawCentredString(
            4.25 * inch, 0.45 * inch,
            f"Decision-support prototype | Page {doc.page}",
        )
        canvas.restoreState()

    def data_table(headers, rows, widths, font_size=7.6):
        data = [[Paragraph(f"<b>{value}</b>", body) for value in headers]]
        for row in rows:
            data.append([Paragraph(str(value), body) for value in row])
        table = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF1F4")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#081F30")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#A8BBC6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ]))
        return table

    def callout(title, text, fill="#E7F5F2"):
        content = Paragraph(f"<b>{title}</b><br/>{text}", body)
        table = Table([[content]], colWidths=[6.5 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(fill)),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#8CA7B5")),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        return KeepTogether([table, Spacer(1, 8)])

    story = [
        Spacer(1, 1.4 * inch),
        Paragraph("TECHNICAL REPORT", ParagraphStyle(
            "Kicker", parent=body, fontName="Helvetica-Bold", fontSize=10,
            alignment=TA_CENTER, textColor=colors.HexColor("#1C9789"), spaceAfter=12,
        )),
        Paragraph("VECTRA-X", cover_title),
        Paragraph(
            "Uncertainty-Aware Triage Intelligence for Vector-Borne Disease Response "
            "in Resource-Limited Humanitarian Settings",
            cover_subtitle,
        ),
        Paragraph(
            "FIT Competition 2026 | Track IV: AI-based Vector-Borne Disease Prediction<br/>"
            "Evidence generated from a self-contained notebook and the official CSV",
            ParagraphStyle("Meta", parent=body, alignment=TA_CENTER, textColor=colors.HexColor("#596771")),
        ),
        PageBreak(),
        Paragraph("Executive Summary", h1),
        Paragraph(
            "VECTRA-X reframes vector-borne disease prediction as an operational triage "
            "problem. It preserves co-infection, separates feature timing, exposes target "
            "leakage, communicates uncertainty, and creates transparent review and resource queues.",
            body,
        ),
        callout(
            "Primary finding",
            f"{manifest['selected_model']} ranked first by training-only OOF macro PR-AUC. "
            f"Calibrated PRE_LAB held-out macro F1 was {heldout['macro_f1']:.3f}, "
            f"micro F1 {heldout['micro_f1']:.3f}, and macro PR-AUC {heldout['macro_pr_auc']:.3f}.",
        ),
        callout(
            "Safety limitation",
            "Yellow fever recall was 0/3 and typhoid recall was 1/7. The system is "
            "decision-support, not diagnosis, and requires prospective validation.",
            "#FBECEB",
        ),
        Paragraph("1. Humanitarian Problem and Intended Use", h1),
        Paragraph(
            "Outbreak teams must decide who needs urgent review, who should receive scarce "
            "confirmatory tests, and when model uncertainty is too high for automated priority. "
            "VECTRA-X supports these decisions without replacing clinician assessment or confirmatory testing.",
            body,
        ),
        Paragraph("2. Related Work and Research Gap", h1),
        Paragraph(
            "Most febrile-illness models optimize one disease or mutually exclusive classes. "
            "VECTRA-X addresses stage-aware, leakage-audited, multi-label triage with uncertainty "
            "and scarce-resource prioritization.",
            body,
        ),
        Paragraph("3. Data and Problem Formulation", h1),
        data_table(
            ["Property", "Observed value", "Consequence"],
            [
                ["Patients", "300", "Use uncertainty intervals"],
                ["Variables", "109", "Audit schema and timing"],
                ["Health centers", "2", "Stress-test facility transfer"],
                ["Multi-label patients", "158", "Use multi-label learning"],
                ["Active labels", "5", "Report macro and per-label metrics"],
            ],
            [1.35 * inch, 1.15 * inch, 4.0 * inch],
        ),
        Spacer(1, 8),
        Paragraph("4. Leakage-Safe Methodology", h1),
        Paragraph(
            "The untouched multi-label holdout is created immediately after target extraction. "
            "All learned schema, imputation, model selection, thresholds, calibration, and "
            "conformal thresholds use training data only.",
            body,
        ),
    ]
    for item in [
        "PRE_LAB: demographics, symptoms, history, and available vital signs.",
        "LAB_AWARE: adds ordered laboratory and rapid-test information.",
        "FULL: includes target-restatement risks and is a research-only leakage demonstration.",
    ]:
        story.append(Paragraph("• " + item, bullet))

    story.extend([
        Paragraph("5. Model Comparison", h1),
        Image(str(model_figure), width=6.2 * inch, height=3.4 * inch),
        Paragraph("Figure 1. Training-only OOF model ranking by macro PR-AUC.", caption),
        data_table(
            ["Model", "OOF PR-AUC", "Runtime"],
            [[r["model"], f"{r['macro_pr_auc']:.3f}", f"{r['runtime_seconds']:.1f}s"] for _, r in models.iterrows()],
            [3.5 * inch, 1.5 * inch, 1.5 * inch],
        ),
        Spacer(1, 8),
        Paragraph(
            "Macro PR-AUC ranks candidates under imbalance. Final triage suitability also "
            "considers recall, false negatives, calibration, stability, and workload.",
            body,
        ),
        Paragraph("6. Held-Out Results", h1),
        data_table(
            ["Track", "Macro F1", "Micro F1", "Macro recall", "Macro PR-AUC"],
            [[r["track"], f"{r['macro_f1']:.3f}", f"{r['micro_f1']:.3f}", f"{r['macro_recall']:.3f}", f"{r['macro_pr_auc']:.3f}"] for _, r in tracks.iterrows()],
            [1.7 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch],
        ),
        Spacer(1, 8),
        Paragraph(
            "FULL performance is higher but operationally invalid because it includes "
            "post-diagnosis signal. PRE_LAB remains the honest early-triage track.",
            body,
        ),
        Image(str(label_figure), width=6.2 * inch, height=3.35 * inch),
        Paragraph("Figure 2. Held-out recall with positive support counts.", caption),
        data_table(
            ["Label", "Support", "Precision", "Recall", "F1", "FNR"],
            [[r["label"], int(r["support_pos"]), f"{r['precision']:.3f}", f"{r['recall']:.3f}", f"{r['f1']:.3f}", f"{r['false_negative_rate']:.3f}"] for _, r in per_label.iterrows()],
            [1.9 * inch, 0.8 * inch, 0.95 * inch, 0.95 * inch, 0.95 * inch, 0.95 * inch],
        ),
        Spacer(1, 8),
        Paragraph("7. Metrics for Real-World Triage", h1),
        data_table(
            ["Question", "Metric", "Operational purpose"],
            [
                ["Rank rare diseases?", "Macro PR-AUC", "Limits common-label dominance"],
                ["Which labels are missed?", "Recall and FNR", "Exposes false negatives"],
                ["How much review work?", "Precision and NNR", "Quantifies staff burden"],
                ["Can probabilities guide priority?", "Brier and ECE", "Tests probability trust"],
                ["Does testing policy help?", "Decision-curve net benefit", "Compares model and default strategies"],
                ["Will evidence transfer?", "CI, repeated CV, center transfer", "Exposes instability"],
            ],
            [2.0 * inch, 1.75 * inch, 2.75 * inch],
        ),
        Spacer(1, 8),
        Paragraph("8. Calibration, Uncertainty, and Caution Sets", h1),
        data_table(
            ["Label", "Support", "Brier raw", "Brier calibrated", "ECE raw", "ECE calibrated"],
            [[r["label"], int(r["support_pos"]), f"{r['brier_raw']:.3f}", f"{r['brier_calibrated']:.3f}", f"{r['ece_raw']:.3f}", f"{r['ece_calibrated']:.3f}"] for _, r in calibration.iterrows()],
            [1.65 * inch, 0.7 * inch, 1.0 * inch, 1.15 * inch, 0.9 * inch, 1.1 * inch],
        ),
        Spacer(1, 8),
        Paragraph(
            "Calibration is label-dependent. Entropy and probability margin route uncertain "
            "cases to human review. Label-wise conformal sets are a caution mechanism, not a "
            "joint multi-label guarantee.",
            body,
        ),
        data_table(
            ["Label", "Test positives", "Coverage", "Inclusions"],
            [[r["label"], int(r["test_positives"]), f"{r['empirical_positive_coverage']:.3f}", int(r["predicted_inclusions"])] for _, r in conformal.iterrows()],
            [2.5 * inch, 1.3 * inch, 1.3 * inch, 1.4 * inch],
        ),
        Spacer(1, 8),
        Paragraph("9. Fairness and Center Robustness", h1),
        Paragraph(
            "Subgroup results are warning diagnostics, not proof of fairness. Facility "
            "transfer is a stronger stress test and shows that prospective deployment "
            "requires center-specific validation and recalibration.",
            body,
        ),
    ])
    if not transfer.empty:
        cols = [c for c in transfer.columns if c in {"test_center", "n_test", "macro_f1", "recall_dengue", "recall_typhoid", "recall_yellow_fever"}]
        story.extend([
            data_table(
                [c.replace("_", " ").title() for c in cols],
                [[f"{v:.3f}" if isinstance(v, float) else v for v in r[cols].tolist()] for _, r in transfer.iterrows()],
                [6.5 * inch / len(cols)] * len(cols),
            ),
            Spacer(1, 8),
        ])
    story.extend([
        Paragraph("10. Triage and Resource Workflow", h1),
        Paragraph(
            "Each case receives separate prediction, uncertainty, triage priority, and "
            "recommended human action fields. Resource simulation makes unserved demand visible.",
            body,
        ),
        data_table(
            ["Resource", "Demand", "Capacity", "Unserved"],
            [[r["resource"], int(r["demand"]), int(r["capacity"]), int(r["unserved"])] for _, r in resources.iterrows()],
            [2.7 * inch, 1.25 * inch, 1.25 * inch, 1.3 * inch],
        ),
        Spacer(1, 8),
        data_table(
            ["Policy", "Flags", "True-positive flags", "Number needed to review", "Flags/patient"],
            [[r["policy"], int(r["total_flags"]), int(r["true_positive_flags"]), f"{r['number_needed_to_review']:.2f}", f"{r['avg_flags_per_patient']:.2f}"] for _, r in thresholds.iterrows()],
            [1.4 * inch, 0.85 * inch, 1.4 * inch, 1.75 * inch, 1.1 * inch],
        ),
        Spacer(1, 8),
        Paragraph("11. Ethics, Privacy, and Limitations", h1),
    ])
    for item in [
        "No treatment recommendation is produced, and clinician assessment remains authoritative.",
        "Rare-label estimates are unstable, especially yellow fever and typhoid.",
        "Public exports use case IDs and exclude UUIDs and truth labels.",
        "Fairness estimates are warning diagnostics, not evidence of equity.",
        "Triage weights and resource policies require clinician co-design.",
    ]:
        story.append(Paragraph("• " + item, bullet))
    story.extend([
        Paragraph("12. Conclusion", h1),
        Paragraph(
            "The strongest contribution is the workflow around an honest pre-lab model: "
            "multi-label framing, leakage control, uncertainty-aware human review, and "
            "transparent resource prioritization. Future work requires prospective temporal "
            "and multi-center validation, larger rare-label cohorts, facility recalibration, "
            "clinician usability testing, and outcome-based resource optimization.",
            body,
        ),
        PageBreak(),
        Paragraph("Appendix A. Confidence Intervals", h1),
        data_table(
            ["Metric", "Estimate", "95% CI low", "95% CI high", "Draws"],
            [[r["metric"], f"{r['estimate']:.3f}", f"{r['ci_low']:.3f}", f"{r['ci_high']:.3f}", int(r["bootstrap_draws"])] for _, r in confidence.iterrows()],
            [2.2 * inch, 1.1 * inch, 1.1 * inch, 1.1 * inch, 1.0 * inch],
        ),
        Spacer(1, 10),
        Paragraph("Appendix B. Reproducibility", h1),
        Paragraph(
            f"Python {manifest['python']}; random state {manifest['random_state']}; "
            f"CSV SHA-256 {manifest['data']['sha256']}. The notebook was executed in a "
            "clean folder containing only the notebook and CSV.",
            body,
        ),
        Paragraph("References", h1),
    ])
    refs = [
        "World Health Organization. Global vector control response 2017-2030. 2017.",
        "World Health Organization. Ethics and governance of artificial intelligence for health. 2021.",
        "World Health Organization. Dengue guidelines. 2009.",
        "World Health Organization. WHO guidelines for malaria.",
        "Sechidis K et al. Multi-label stratification. ECML PKDD, 2011.",
        "Read J et al. Classifier chains. Machine Learning, 2011.",
        "Niculescu-Mizil A, Caruana R. Probability calibration. ICML, 2005.",
        "Angelopoulos AN, Bates S. Conformal prediction. 2023.",
        "Vickers AJ, Elkin EB. Decision curve analysis. 2006.",
        "Collins GS et al. TRIPOD+AI. BMJ, 2024.",
        "Wolff RF et al. PROBAST. 2019.",
        "Lundberg SM, Lee SI. SHAP. NeurIPS, 2017.",
        "Chen T, Guestrin C. XGBoost. KDD, 2016.",
        "Ke G et al. LightGBM. NeurIPS, 2017.",
    ]
    for index, ref in enumerate(refs, 1):
        story.append(Paragraph(f"{index}. {ref}", body))

    document = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=LETTER,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.82 * inch,
        bottomMargin=0.7 * inch,
        title="VECTRA-X Technical Report",
        author="VECTRA-X Team",
    )
    document.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((DATA / "execution_manifest.json").read_text(encoding="utf-8"))
    models = load("model_comparison.csv")
    rationale = load("model_rationale.csv")
    heldout = load("held_out_metrics.csv").iloc[0]
    per_label = load("per_label_metrics.csv")
    tracks = load("held_out_track_comparison.csv")
    confidence = load("confidence_intervals.csv")
    calibration = load("calibration_metrics.csv")
    conformal = load("conformal_metrics.csv")
    transfer = load("center_transfer.csv")
    resources = load("resource_simulation.csv")
    thresholds = load("threshold_resource_tradeoff.csv")
    model_figure, label_figure = make_figures(models, per_label)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, before, after in [
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE if style_name != "Heading 3" else NAVY
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("VECTRA-X | FIT Competition 2026, Track IV"), size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        footer.add_run("Decision-support prototype | Requires prospective validation"),
        size=8.5,
        color=GRAY,
    )

    # Editorial cover.
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("TECHNICAL REPORT"), size=11, bold=True, color=TEAL)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_font(title.add_run("VECTRA-X"), size=30, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    set_font(
        subtitle.add_run(
            "Uncertainty-Aware Triage Intelligence for Vector-Borne Disease Response "
            "in Resource-Limited Humanitarian Settings"
        ),
        size=16,
        color=BLUE,
    )
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        metadata.add_run(
            "FIT Competition 2026 | Track IV: AI-based Vector-Borne Disease Prediction\n"
            "Evidence generated from a self-contained notebook and the official CSV"
        ),
        size=10.5,
        color=GRAY,
        italic=True,
    )
    doc.add_page_break()

    add_heading(doc, "Executive Summary", 1)
    add_paragraph(
        doc,
        "VECTRA-X reframes vector-borne disease prediction as an operational triage "
        "problem. The system preserves co-infection, separates pre-lab and lab-aware "
        "information, exposes target leakage, communicates uncertainty, and converts "
        "model outputs into transparent review and resource-prioritization queues."
    )
    add_callout(
        doc,
        "Primary finding",
        f"{manifest['selected_model']} ranked first by training-only OOF macro PR-AUC. "
        f"On the untouched holdout, calibrated PRE_LAB performance reached macro F1 "
        f"{heldout['macro_f1']:.3f}, micro F1 {heldout['micro_f1']:.3f}, and macro "
        f"PR-AUC {heldout['macro_pr_auc']:.3f}. These aggregate scores conceal major "
        "rare-label false-negative risk.",
    )
    add_callout(
        doc,
        "Safety limitation",
        "Yellow fever recall was 0/3 and typhoid recall was 1/7. VECTRA-X is therefore "
        "a competition decision-support prototype, not a diagnostic system, and "
        "requires prospective multi-center validation.",
        fill=PALE_RED,
    )

    add_heading(doc, "1. Humanitarian Problem and Intended Use", 1)
    add_paragraph(
        doc,
        "During outbreaks, frontline teams must prioritize patients while laboratory "
        "tests, beds, and qualified staff are limited. A prediction alone does not "
        "answer who should be reviewed first, when the model is uncertain, or how "
        "capacity constraints change the queue."
    )
    add_paragraph(
        doc,
        "VECTRA-X supports triage prioritization and confirmatory-test allocation. "
        "Its intended users are health workers and humanitarian response coordinators. "
        "It does not recommend treatment or replace clinician assessment."
    )

    add_heading(doc, "2. Related Work and Research Gap", 1)
    add_paragraph(
        doc,
        "Clinical prediction studies for febrile and vector-borne illness often focus "
        "on one disease or force mutually exclusive classes. This is poorly matched to "
        "a cohort in which 158 of 300 patients have more than one active label."
    )
    add_paragraph(
        doc,
        "The research gap addressed here is stage-aware, leakage-audited, multi-label "
        "triage with uncertainty and operational prioritization. The methodology draws "
        "on multi-label stratification, probability calibration, conformal prediction, "
        "decision-curve analysis, TRIPOD+AI reporting, and WHO guidance on responsible "
        "AI for health."
    )

    add_heading(doc, "3. Data and Problem Formulation", 1)
    add_table(
        doc,
        ["Property", "Observed value", "Methodological consequence"],
        [
            ["Patients", "300", "Use uncertainty intervals and repeated validation"],
            ["Raw variables", "109", "Audit schema, missingness, and feature timing"],
            ["Health centers", "2", "Run leave-one-center-out stress tests"],
            ["Multi-label patients", "158", "Use multi-label rather than multi-class learning"],
            ["Active labels", "5", "Report macro and per-label metrics"],
            ["Dataset checksum", manifest["data"]["sha256"][:16] + "...", "Trace evidence to the submitted CSV"],
        ],
        [2100, 1800, 5460],
    )

    add_heading(doc, "4. Leakage-Safe Methodology", 1)
    add_paragraph(
        doc,
        "The outer multi-label stratified holdout is created immediately after target "
        "extraction. Data-driven schema decisions, category discovery, missing-indicator "
        "selection, imputation, model selection, threshold optimization, calibration, "
        "and conformal thresholds use training data only."
    )
    add_bullet(doc, "PRE_LAB uses demographics, symptoms, history, and available vital signs.")
    add_bullet(doc, "LAB_AWARE adds ordered laboratory and rapid-test information.")
    add_bullet(doc, "FULL includes target-restatement risks and is reported only as a leakage demonstration.")
    add_paragraph(
        doc,
        "All notebook results are recomputed from the official CSV. The submitted "
        "notebook does not read project modules, configuration files, processed data, "
        "saved models, or prior artifacts."
    )

    add_heading(doc, "5. Model Comparison and Selection Rationale", 1)
    doc.add_picture(str(model_figure), width=Inches(6.35))
    caption = doc.add_paragraph("Figure 1. Training-only OOF model ranking by macro PR-AUC.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(caption.runs[0], size=9, italic=True, color=GRAY)
    model_rows = []
    rationale_map = rationale.set_index("model")
    for _, row in models.iterrows():
        detail = rationale_map.loc[row["model"]]
        model_rows.append([
            row["model"],
            f"{row['macro_pr_auc']:.3f}",
            f"{row['runtime_seconds']:.1f}s",
            detail["rationale"],
            detail["limitation"],
        ])
    add_table(
        doc,
        ["Model", "OOF PR-AUC", "Runtime", "Why included", "Main limitation"],
        model_rows,
        [1450, 1050, 800, 3000, 3060],
    )
    add_paragraph(
        doc,
        "Macro PR-AUC is the initial selection metric because label prevalence is highly "
        "imbalanced and the metric evaluates ranking without fixing a threshold. Final "
        "triage suitability also considers recall, false negatives, calibration, "
        "stability, number-needed-to-review, and resource demand."
    )

    add_heading(doc, "6. Held-Out Results", 1)
    add_table(
        doc,
        ["Track", "Macro F1", "Micro F1", "Macro recall", "Macro PR-AUC"],
        [
            [
                row["track"],
                f"{row['macro_f1']:.3f}",
                f"{row['micro_f1']:.3f}",
                f"{row['macro_recall']:.3f}",
                f"{row['macro_pr_auc']:.3f}",
            ]
            for _, row in tracks.iterrows()
        ],
        [2300, 1765, 1765, 1765, 1765],
    )
    add_paragraph(
        doc,
        "The FULL track is highest, but its gain is not operationally valid because it "
        "contains post-diagnosis or target-restatement information. LAB_AWARE is a "
        "confirmation-support context. PRE_LAB remains the honest early-triage track."
    )
    doc.add_picture(str(label_figure), width=Inches(6.35))
    caption = doc.add_paragraph("Figure 2. Held-out recall with positive support counts.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(caption.runs[0], size=9, italic=True, color=GRAY)
    add_table(
        doc,
        ["Label", "Support", "Precision", "Recall", "F1", "FNR"],
        [
            [
                row["label"],
                int(row["support_pos"]),
                f"{row['precision']:.3f}",
                f"{row['recall']:.3f}",
                f"{row['f1']:.3f}",
                f"{row['false_negative_rate']:.3f}",
            ]
            for _, row in per_label.iterrows()
        ],
        [2400, 1100, 1465, 1465, 1465, 1465],
    )

    add_heading(doc, "7. Metric Rationale for Real-World Triage", 1)
    add_table(
        doc,
        ["Question", "Metric", "Why it matters"],
        [
            ["Can the model rank rare diseases?", "Macro PR-AUC", "Resists dominance by common labels"],
            ["Which diseases are missed?", "Per-label recall and FNR", "Makes false negatives visible"],
            ["How much review work is created?", "Precision and number-needed-to-review", "Connects flags to staff burden"],
            ["Can probability guide priority?", "Brier score and calibration error", "Tests probability trust"],
            ["Is testing policy useful?", "Decision-curve net benefit", "Compares model, test-all, and test-none"],
            ["Will evidence transfer?", "Repeated CV, bootstrap CI, center transfer", "Exposes sampling and facility instability"],
        ],
        [2600, 2600, 4160],
    )
    add_paragraph(
        doc,
        "No single metric is treated as a clinical utility score. The recommended "
        "candidate is reviewed as a Pareto trade-off across discrimination, safety, "
        "probability trust, stability, and operational burden."
    )

    add_heading(doc, "8. Calibration, Uncertainty, and Caution Sets", 1)
    add_table(
        doc,
        ["Label", "Support", "Brier raw", "Brier calibrated", "ECE raw", "ECE calibrated"],
        [
            [
                row["label"], int(row["support_pos"]),
                f"{row['brier_raw']:.3f}", f"{row['brier_calibrated']:.3f}",
                f"{row['ece_raw']:.3f}", f"{row['ece_calibrated']:.3f}",
            ]
            for _, row in calibration.iterrows()
        ],
        [2200, 1000, 1540, 1540, 1540, 1540],
    )
    add_paragraph(
        doc,
        "Calibration is label-dependent and may worsen some estimates. Predictive "
        "entropy, top-two margin, and maximum probability form an uncertainty gate. "
        "Label-wise positive-class conformal sets are presented as a caution mechanism, "
        "not as a joint multi-label coverage guarantee."
    )
    add_table(
        doc,
        ["Label", "Test positives", "Coverage", "Inclusions"],
        [
            [
                row["label"],
                int(row["test_positives"]),
                f"{row['empirical_positive_coverage']:.3f}",
                int(row["predicted_inclusions"]),
            ]
            for _, row in conformal.iterrows()
        ],
        [3300, 1800, 2100, 2160],
    )

    add_heading(doc, "9. Explainability, Fairness, and Robustness", 1)
    add_paragraph(
        doc,
        "Permutation importance explains model behavior globally. It does not establish "
        "medical causality. Subgroup metrics are warning diagnostics because support is "
        "small. Center transfer is a stronger stress test of deployment risk."
    )
    if not transfer.empty:
        transfer_columns = [c for c in transfer.columns if c in {"test_center", "n_test", "macro_f1", "recall_dengue", "recall_typhoid", "recall_yellow_fever"}]
        add_table(
            doc,
            [column.replace("_", " ").title() for column in transfer_columns],
            [
                [
                    f"{value:.3f}" if isinstance(value, float) else str(value)
                    for value in row[transfer_columns].tolist()
                ]
                for _, row in transfer.iterrows()
            ],
            [1700, 1100, 1300, 1750, 1750, 1760][: len(transfer_columns)],
        )
    add_callout(
        doc,
        "Generalization warning",
        "Performance across facilities is materially weaker for rare labels. A future "
        "deployment would require facility-specific validation, recalibration, drift "
        "monitoring, and a human override pathway.",
        fill=PALE_RED,
    )

    add_heading(doc, "10. Triage and Resource Workflow", 1)
    add_paragraph(
        doc,
        "The prototype separates four outputs: disease predictions, uncertainty, triage "
        "priority, and recommended human action. Tiers are Routine, Clinical Review, "
        "Confirmatory Test Priority, and Urgent Response."
    )
    add_table(
        doc,
        ["Resource", "Demand", "Capacity", "Unserved"],
        [
            [row["resource"], int(row["demand"]), int(row["capacity"]), int(row["unserved"])]
            for _, row in resources.iterrows()
        ],
        [3500, 1950, 1950, 1960],
    )
    add_table(
        doc,
        ["Policy", "Total flags", "True-positive flags", "Number needed to review", "Flags per patient"],
        [
            [
                row["policy"], int(row["total_flags"]), int(row["true_positive_flags"]),
                f"{row['number_needed_to_review']:.2f}", f"{row['avg_flags_per_patient']:.2f}",
            ]
            for _, row in thresholds.iterrows()
        ],
        [1850, 1650, 1900, 2300, 1660],
    )
    add_paragraph(
        doc,
        "These are workflow simulations, not observed clinical outcomes. Their purpose "
        "is to make the sensitivity-workload trade-off visible to decision makers."
    )

    add_heading(doc, "11. Ethics, Privacy, and Limitations", 1)
    for statement in [
        "VECTRA-X does not replace clinician assessment or provide treatment recommendations.",
        "Rare-label estimates are unstable, especially yellow fever and typhoid.",
        "Public dashboard exports use generated case IDs and exclude UUIDs and truth labels.",
        "Fairness estimates are warning diagnostics, not evidence of equity.",
        "Conformal coverage is label-wise and approximate under small calibration samples.",
        "Triage weights and resource policies require clinician and humanitarian-operations co-design.",
    ]:
        add_bullet(doc, statement)

    add_heading(doc, "12. Conclusion and Recommendations", 1)
    add_paragraph(
        doc,
        "VECTRA-X demonstrates a defensible humanitarian workflow around an honest "
        "pre-lab model. Its strongest contribution is not an aggregate score; it is the "
        "combination of multi-label framing, leakage control, uncertainty-aware human "
        "review, and transparent resource prioritization."
    )
    add_paragraph(
        doc,
        "Before operational use, the next steps are prospective temporal and multi-center "
        "validation, larger rare-label cohorts, facility recalibration, clinician usability "
        "testing, and outcome-based resource optimization."
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix A. Confidence Intervals", 1)
    add_table(
        doc,
        ["Metric", "Estimate", "95% CI low", "95% CI high", "Bootstrap draws"],
        [
            [
                row["metric"], f"{row['estimate']:.3f}", f"{row['ci_low']:.3f}",
                f"{row['ci_high']:.3f}", int(row["bootstrap_draws"]),
            ]
            for _, row in confidence.iterrows()
        ],
        [2800, 1640, 1640, 1640, 1640],
    )

    add_heading(doc, "Appendix B. Reproducibility", 1)
    add_paragraph(
        doc,
        f"The executed notebook used Python {manifest['python']}, random state "
        f"{manifest['random_state']}, and CSV SHA-256 {manifest['data']['sha256']}. "
        "It was verified in a clean folder containing only the notebook and CSV."
    )
    add_paragraph(
        doc,
        "The notebook exports fresh model comparison, held-out metrics, per-label "
        "metrics, confidence intervals, calibration, conformal, fairness, center "
        "transfer, resource trade-offs, privacy-safe patient records, and an execution manifest."
    )

    add_heading(doc, "References", 1)
    references = [
        "World Health Organization. Global vector control response 2017-2030. 2017.",
        "World Health Organization. Ethics and governance of artificial intelligence for health. 2021.",
        "World Health Organization. Dengue: guidelines for diagnosis, treatment, prevention and control. 2009.",
        "World Health Organization. WHO guidelines for malaria. Continuously updated.",
        "Sechidis K, Tsoumakas G, Vlahavas I. On the stratification of multi-label data. ECML PKDD, 2011.",
        "Read J, Pfahringer B, Holmes G, Frank E. Classifier chains for multi-label classification. Machine Learning, 2011.",
        "Niculescu-Mizil A, Caruana R. Predicting good probabilities with supervised learning. ICML, 2005.",
        "Vovk V, Gammerman A, Shafer G. Algorithmic Learning in a Random World. Springer, 2005.",
        "Angelopoulos AN, Bates S. Conformal prediction: a gentle introduction. Foundations and Trends in Machine Learning, 2023.",
        "Vickers AJ, Elkin EB. Decision curve analysis. Medical Decision Making, 2006.",
        "Collins GS et al. TRIPOD+AI statement. BMJ, 2024.",
        "Wolff RF et al. PROBAST. Annals of Internal Medicine, 2019.",
        "Lundberg SM, Lee SI. A unified approach to interpreting model predictions. NeurIPS, 2017.",
        "Chen T, Guestrin C. XGBoost. KDD, 2016.",
        "Ke G et al. LightGBM. NeurIPS, 2017.",
        "Obermeyer Z et al. Dissecting racial bias in a population health algorithm. Science, 2019.",
    ]
    for index, reference in enumerate(references, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"{index}. {reference}")
        set_font(run, size=9.5)

    doc.core_properties.title = "VECTRA-X Technical Report"
    doc.core_properties.subject = "FIT Competition 2026 Track IV"
    doc.core_properties.author = "VECTRA-X Team"
    doc.save(DOCX_PATH)
    build_pdf(
        manifest, models, heldout, per_label, tracks, confidence, calibration,
        conformal, transfer, resources, thresholds, model_figure, label_figure,
    )
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    build()
